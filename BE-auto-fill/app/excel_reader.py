"""Read an uploaded Excel workbook into a header + list of rows."""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

import openpyxl

SPECIAL_COLUMNS = {"pagehistory": "pageHistory", "emailaddress": "emailAddress"}


class ExcelError(Exception):
    """Raised when the uploaded workbook cannot be read."""


@dataclass
class ExcelData:
    sheet_name: str
    sheet_names: list[str]
    header: list[str]
    rows: list[list[Any]] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return len(self.rows)


def normalize_cell(val: Any) -> str | None:
    """Convert a cell value to the string Google Forms expects, or ``None`` if empty."""
    if val is None:
        return None
    if isinstance(val, bool):
        return "TRUE" if val else "FALSE"
    if isinstance(val, float):
        if val != val:  # NaN
            return None
        if val.is_integer():
            return str(int(val))
        return str(val)
    if isinstance(val, int):
        return str(val)
    if isinstance(val, dt.datetime):
        if val.time() == dt.time(0, 0):
            return val.strftime("%Y-%m-%d")
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, dt.date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, dt.time):
        return val.strftime("%H:%M")
    text = str(val).strip()
    if text == "" or text.lower() in {"nan", "none", "null"}:
        return None
    return text


def _header_name(val: Any, idx: int) -> str:
    text = normalize_cell(val)
    if text is None:
        return f"col_{idx}"
    return text


def read_excel(content: bytes, sheet_name: str | None = None, header_row: int = 1) -> ExcelData:
    """Load ``content`` (xlsx bytes) and return the chosen sheet as header + rows.

    * ``sheet_name``: defaults to a sheet called ``Coded`` if present, otherwise
      the first sheet.
    * ``header_row``: 1-based index of the header row. Rows above it are ignored.
    """
    if not content:
        raise ExcelError("File Excel rỗng.")
    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises many different error types
        raise ExcelError(f"Không đọc được file Excel (.xlsx): {exc}") from exc

    sheet_names = list(wb.sheetnames)
    if not sheet_names:
        raise ExcelError("Workbook không có sheet nào.")

    if sheet_name:
        if sheet_name not in sheet_names:
            raise ExcelError(
                f"Không tìm thấy sheet '{sheet_name}'. Các sheet hiện có: {', '.join(sheet_names)}"
            )
        chosen = sheet_name
    elif "Coded" in sheet_names:
        chosen = "Coded"
    else:
        chosen = sheet_names[0]

    ws = wb[chosen]
    if header_row < 1:
        raise ExcelError("header_row phải >= 1.")

    all_rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    wb.close()
    if not all_rows:
        raise ExcelError(f"Sheet '{chosen}' không có dữ liệu.")

    raw_header = list(all_rows[0])
    # Trim trailing empty header cells
    while raw_header and normalize_cell(raw_header[-1]) is None:
        raw_header.pop()
    if not raw_header:
        raise ExcelError(f"Dòng tiêu đề (dòng {header_row}) của sheet '{chosen}' trống.")
    header = [_header_name(v, i) for i, v in enumerate(raw_header)]

    rows: list[list[Any]] = []
    for r in all_rows[1:]:
        values = list(r)[: len(header)]
        if all(normalize_cell(v) is None for v in values):
            continue  # skip fully empty rows
        values.extend([None] * (len(header) - len(values)))
        rows.append(values)

    return ExcelData(sheet_name=chosen, sheet_names=sheet_names, header=header, rows=rows)


def list_sheets(content: bytes) -> list[str]:
    try:
        wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    except Exception as exc:
        raise ExcelError(f"Không đọc được file Excel (.xlsx): {exc}") from exc
    names = list(wb.sheetnames)
    wb.close()
    return names
